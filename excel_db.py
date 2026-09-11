from openpyxl import Workbook, load_workbook
import os
from datetime import datetime
import hashlib
import random

class ExcelDB:
    def __init__(self, data_dir='data'):
        self.data_dir = data_dir
        self.users_file = os.path.join(data_dir, 'users.xlsx')
        self.events_file = os.path.join(data_dir, 'events.xlsx')
        self.bookings_file = os.path.join(data_dir, 'bookings.xlsx')
        self.venues_file = os.path.join(data_dir, 'venues.xlsx')
        self.feedback_file = os.path.join(data_dir, 'feedback.xlsx')
        self.init_files()
    
    def init_files(self):
        if not os.path.exists(self.data_dir):
            os.makedirs(self.data_dir)
        
        # Initialize Users
        if not os.path.exists(self.users_file):
            wb = Workbook()
            ws = wb.active
            ws.append(['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
            admin_hash = hashlib.sha256('admin123'.encode()).hexdigest()
            organizer_hash = hashlib.sha256('organizer123'.encode()).hexdigest()
            user_hash = hashlib.sha256('user123'.encode()).hexdigest()
            ws.append([1, 'Admin User', 'admin@example.com', 'admin', admin_hash, 0, False])
            ws.append([2, 'Event Organizer', 'organizer@example.com', 'organizer', organizer_hash, 0, False])
            ws.append([3, 'Regular User', 'user@example.com', 'user', user_hash, 0, False])
            wb.save(self.users_file)
        
        # Initialize Events
        if not os.path.exists(self.events_file):
            wb = Workbook()
            ws = wb.active
            ws.append(['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
            ws.append([1, 'Rock Concert', 'Music', '2024-12-25', 'Stadium A', 2, 1000, 950, 50.0])
            ws.append([2, 'Comedy Night', 'Comedy', '2024-12-30', 'Theater B', 2, 500, 480, 25.0])
            ws.append([3, 'Tech Conference', 'Conference', '2025-01-15', 'Convention Center', 2, 200, 180, 100.0])
            wb.save(self.events_file)
        
        # Initialize Bookings
        if not os.path.exists(self.bookings_file):
            wb = Workbook()
            ws = wb.active
            ws.append(['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status'])
            wb.save(self.bookings_file)
        
        # Initialize Venues
        if not os.path.exists(self.venues_file):
            wb = Workbook()
            ws = wb.active
            ws.append(['VenueID', 'Name', 'City', 'Capacity'])
            ws.append([1, 'Stadium A', 'New York', 1000])
            ws.append([2, 'Theater B', 'Los Angeles', 500])
            ws.append([3, 'Convention Center', 'Chicago', 200])
            wb.save(self.venues_file)
        
        # Initialize Feedback
        if not os.path.exists(self.feedback_file):
            wb = Workbook()
            ws = wb.active
            ws.append(['FeedbackID', 'UserID', 'EventID', 'Rating', 'Comments'])
            wb.save(self.feedback_file)
    
    def read_excel(self, file_path):
        try:
            wb = load_workbook(file_path)
            ws = wb.active
            data = []
            headers = [cell.value for cell in ws[1] if cell.value]
            if not headers:
                return []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(row):
                    row_dict = {}
                    for i, header in enumerate(headers):
                        if i < len(row):
                            row_dict[header] = row[i]
                    if row_dict:
                        data.append(row_dict)
            return data
        except FileNotFoundError:
            return []
        except Exception as e:
            print(f"Error reading Excel file {file_path}: {str(e)}")
            return []
    
    def write_excel(self, data, file_path, headers):
        try:
            wb = Workbook()
            ws = wb.active
            ws.append(headers)
            for row in data:
                if row:
                    ws.append([row.get(h) for h in headers])
            wb.save(file_path)
            return True
        except PermissionError:
            print(f"Permission denied writing to {file_path}. File may be open in Excel.")
            return False
        except Exception as e:
            print(f"Error writing Excel file {file_path}: {str(e)}")
            return False
    
    def get_users(self):
        return self.read_excel(self.users_file)
    
    def get_events(self):
        return self.read_excel(self.events_file)
    
    def get_bookings(self):
        return self.read_excel(self.bookings_file)
    
    def get_venues(self):
        return self.read_excel(self.venues_file)
    
    def get_feedback(self):
        return self.read_excel(self.feedback_file)
    
    def add_user(self, name, email, role, password):
        try:
            if not all([name, email, role, password]):
                raise ValueError("All user fields are required")
            
            users = self.get_users()
            
            # Check if email already exists
            if any(u and u.get('Email') == email for u in users):
                raise ValueError("Email already exists")
            
            user_id = len(users) + 1
            password_hash = hashlib.sha256(password.encode()).hexdigest()
            new_user = {'UserID': user_id, 'Name': name, 'Email': email, 'Role': role, 'PasswordHash': password_hash, 'FailedAttempts': 0, 'IsLocked': False}
            users.append(new_user)
            
            if self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked']):
                return user_id
            else:
                raise Exception("Failed to save user data")
        except Exception as e:
            print(f"Error adding user: {str(e)}")
            return None
    
    def authenticate_user(self, email, password):
        try:
            if not email or not password:
                return None
            
            users = self.get_users()
            password_hash = hashlib.sha256(password.encode()).hexdigest()
            
            for user in users:
                if user and user.get('Email') == email and user.get('PasswordHash') == password_hash:
                    return user
            return None
        except Exception as e:
            print(f"Error authenticating user: {str(e)}")
            return None
    
    def add_event(self, title, category, date, venue, organizer_id, total_seats, price):
        events = self.get_events()
        
        # Generate random event ID
        while True:
            event_id = random.randint(1000, 999999)
            # Check if ID already exists
            if not any(e and e.get('EventID') == event_id for e in events):
                break
        
        events.append({
            'EventID': event_id, 'Title': title, 'Category': category, 'Date': date,
            'Venue': venue, 'OrganizerID': organizer_id, 'TotalSeats': total_seats,
            'AvailableSeats': total_seats, 'Price': price
        })
        self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
        return event_id
    
    def book_tickets(self, user_id, event_id, seats_booked):
        events = self.get_events()
        for i, event in enumerate(events):
            if event and event.get('EventID') == event_id:
                if event.get('AvailableSeats', 0) >= seats_booked:
                    events[i]['AvailableSeats'] = event['AvailableSeats'] - seats_booked
                    self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
                    
                    bookings = self.get_bookings()
                    booking_id = len(bookings) + 1
                    bookings.append({
                        'BookingID': booking_id, 'UserID': user_id, 'EventID': event_id,
                        'SeatsBooked': seats_booked, 'BookingDate': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                        'Status': 'Confirmed', 'Amount': 0, 'CardLast4': ''
                    })
                    self.write_excel(bookings, self.bookings_file, ['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4'])
                    return booking_id
        return None
    
    def book_tickets_with_payment(self, user_id, event_id, seats_booked, amount, card_last4):
        try:
            if not all([user_id, event_id, seats_booked, amount]) or seats_booked <= 0:
                return None
            
            events = self.get_events()
            event_found = False
            
            for i, event in enumerate(events):
                if event and event.get('EventID') == event_id:
                    available_seats = event.get('AvailableSeats', 0)
                    if available_seats >= seats_booked:
                        events[i]['AvailableSeats'] = available_seats - seats_booked
                        
                        if not self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price']):
                            return None
                        
                        bookings = self.get_bookings()
                        booking_id = len(bookings) + 1
                        new_booking = {
                            'BookingID': booking_id, 'UserID': user_id, 'EventID': event_id,
                            'SeatsBooked': seats_booked, 'BookingDate': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                            'Status': 'Confirmed', 'Amount': amount, 'CardLast4': card_last4
                        }
                        bookings.append(new_booking)
                        
                        if self.write_excel(bookings, self.bookings_file, ['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4']):
                            return booking_id
                        else:
                            # Rollback seat update
                            events[i]['AvailableSeats'] = available_seats
                            self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
                            return None
                    event_found = True
                    break
            
            return None
        except Exception as e:
            print(f"Error booking tickets: {str(e)}")
            return None
    
    def cancel_booking(self, booking_id, user_id):
        bookings = self.get_bookings()
        events = self.get_events()
        
        for i, booking in enumerate(bookings):
            if booking and booking.get('BookingID') == booking_id and booking.get('UserID') == user_id:
                if booking.get('Status') == 'Confirmed':
                    # Update booking status
                    bookings[i]['Status'] = 'Cancelled'
                    
                    # Restore seats
                    event_id = booking['EventID']
                    seats = booking['SeatsBooked']
                    for j, event in enumerate(events):
                        if event and event.get('EventID') == event_id:
                            events[j]['AvailableSeats'] = event['AvailableSeats'] + seats
                            break
                    
                    # Save changes
                    self.write_excel(bookings, self.bookings_file, ['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4'])
                    self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
                    return True
        return False
    
    def delete_event(self, event_id, organizer_id):
        # Protect default events (IDs 1-3)
        if event_id <= 3:
            return False
        
        events = self.get_events()
        bookings = self.get_bookings()
        
        # Check if event has bookings
        event_bookings = [b for b in bookings if b and b.get('EventID') == event_id and b.get('Status') == 'Confirmed']
        if event_bookings:
            return False  # Cannot delete event with bookings
        
        # Remove event
        events = [e for e in events if not (e and e.get('EventID') == event_id and e.get('OrganizerID') == organizer_id)]
        self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
        return True
    
    def cancel_event(self, event_id, organizer_id):
        # Protect default events (IDs 1-3)
        if event_id <= 3:
            return False
        
        events = self.get_events()
        bookings = self.get_bookings()
        
        # Find event
        event_found = False
        for i, event in enumerate(events):
            if event and event.get('EventID') == event_id and event.get('OrganizerID') == organizer_id:
                events[i]['Status'] = 'Cancelled'
                event_found = True
                break
        
        if not event_found:
            return False
        
        # Cancel all bookings for this event
        for i, booking in enumerate(bookings):
            if booking and booking.get('EventID') == event_id and booking.get('Status') == 'Confirmed':
                bookings[i]['Status'] = 'Cancelled - Event Cancelled'
        
        # Save changes
        self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price', 'Status'])
        self.write_excel(bookings, self.bookings_file, ['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4'])
        return True
    
    def get_user_bookings_with_events(self, user_id):
        bookings = self.get_bookings()
        events = self.get_events()
        
        user_bookings = [b for b in bookings if b and b.get('UserID') == user_id]
        
        # Add event details to bookings
        for booking in user_bookings:
            for event in events:
                if event and event.get('EventID') == booking.get('EventID'):
                    booking['EventTitle'] = event.get('Title')
                    booking['EventDate'] = event.get('Date')
                    booking['EventVenue'] = event.get('Venue')
                    booking['EventCategory'] = event.get('Category')
                    break
        
        return user_bookings
    
    def delete_user(self, user_id, admin_id):
        # Protect default users (IDs 1-3)
        if user_id <= 3:
            return False
        
        # Only admin can delete users
        admin_users = [u for u in self.get_users() if u and u.get('UserID') == admin_id and u.get('Role') == 'admin']
        if not admin_users:
            return False
        
        users = self.get_users()
        users = [u for u in users if not (u and u.get('UserID') == user_id)]
        self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
        return True
    
    def get_event_attendees(self, event_id):
        bookings = self.get_bookings()
        users = self.get_users()
        
        # Get confirmed bookings for this event
        event_bookings = [b for b in bookings if b and b.get('EventID') == event_id and b.get('Status') == 'Confirmed']
        
        # Add user details to bookings
        attendees = []
        for booking in event_bookings:
            for user in users:
                if user and user.get('UserID') == booking.get('UserID'):
                    attendee = {
                        'BookingID': booking.get('BookingID'),
                        'UserName': user.get('Name'),
                        'UserEmail': user.get('Email'),
                        'SeatsBooked': booking.get('SeatsBooked'),
                        'BookingDate': booking.get('BookingDate'),
                        'Amount': booking.get('Amount', 0),
                        'CardLast4': booking.get('CardLast4', '')
                    }
                    attendees.append(attendee)
                    break
        
        return attendees
    
    def update_password(self, email, new_password):
        """Update user password"""
        users = self.get_users()
        password_hash = hashlib.sha256(new_password.encode()).hexdigest()
        
        for i, user in enumerate(users):
            if user and user.get('Email') == email:
                users[i]['PasswordHash'] = password_hash
                self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
                return True
        return False
    
    def update_event(self, event_id, organizer_id, title, category, date, venue, total_seats, price):
        # Protect default events (IDs 1-3)
        if event_id <= 3:
            return False
        
        events = self.get_events()
        for i, event in enumerate(events):
            if event and event.get('EventID') == event_id and event.get('OrganizerID') == organizer_id:
                events[i]['Title'] = title
                events[i]['Category'] = category
                events[i]['Date'] = date
                events[i]['Venue'] = venue
                events[i]['TotalSeats'] = total_seats
                events[i]['Price'] = price
                self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
                return True
        return False
    
    def reset_to_defaults(self):
        # Reset users - keep only default users (IDs 1-3)
        wb = Workbook()
        ws = wb.active
        ws.append(['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
        admin_hash = hashlib.sha256('admin123'.encode()).hexdigest()
        organizer_hash = hashlib.sha256('organizer123'.encode()).hexdigest()
        user_hash = hashlib.sha256('user123'.encode()).hexdigest()
        ws.append([1, 'Admin User', 'admin@example.com', 'admin', admin_hash, 0, False])
        ws.append([2, 'Event Organizer', 'organizer@example.com', 'organizer', organizer_hash, 0, False])
        ws.append([3, 'Regular User', 'user@example.com', 'user', user_hash, 0, False])
        wb.save(self.users_file)
        
        # Reset events - keep only default events (IDs 1-3)
        wb = Workbook()
        ws = wb.active
        ws.append(['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
        ws.append([1, 'Rock Concert', 'Music', '2024-12-25', 'Stadium A', 2, 1000, 950, 50.0])
        ws.append([2, 'Comedy Night', 'Comedy', '2024-12-30', 'Theater B', 2, 500, 480, 25.0])
        ws.append([3, 'Tech Conference', 'Conference', '2025-01-15', 'Convention Center', 2, 200, 180, 100.0])
        wb.save(self.events_file)
        
        # Clear all bookings
        wb = Workbook()
        ws = wb.active
        ws.append(['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4'])
        wb.save(self.bookings_file)
        
        # Reset venues
        wb = Workbook()
        ws = wb.active
        ws.append(['VenueID', 'Name', 'City', 'Capacity'])
        ws.append([1, 'Stadium A', 'New York', 1000])
        ws.append([2, 'Theater B', 'Los Angeles', 500])
        ws.append([3, 'Convention Center', 'Chicago', 200])
        wb.save(self.venues_file)
        
        # Clear feedback
        wb = Workbook()
        ws = wb.active
        ws.append(['FeedbackID', 'UserID', 'EventID', 'Rating', 'Comments'])
        wb.save(self.feedback_file)
    
    def is_account_locked(self, email):
        users = self.get_users()
        for user in users:
            if user and user.get('Email') == email:
                return user.get('IsLocked', False)
        return False
    
    def increment_failed_attempts(self, email):
        users = self.get_users()
        for i, user in enumerate(users):
            if user and user.get('Email') == email:
                users[i]['FailedAttempts'] = user.get('FailedAttempts', 0) + 1
                self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
                return True
        return False
    
    def reset_failed_attempts(self, email):
        users = self.get_users()
        for i, user in enumerate(users):
            if user and user.get('Email') == email:
                users[i]['FailedAttempts'] = 0
                users[i]['IsLocked'] = False
                self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
                return True
        return False
    
    def lock_account(self, email):
        users = self.get_users()
        for i, user in enumerate(users):
            if user and user.get('Email') == email:
                users[i]['IsLocked'] = True
                self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
                return True
        return False
    
    def unlock_account_by_id(self, user_id):
        users = self.get_users()
        for i, user in enumerate(users):
            if user and user.get('UserID') == user_id:
                users[i]['IsLocked'] = False
                users[i]['FailedAttempts'] = 0
                self.write_excel(users, self.users_file, ['UserID', 'Name', 'Email', 'Role', 'PasswordHash', 'FailedAttempts', 'IsLocked'])
                return True
        return False
    
    def admin_delete_event(self, event_id):
        """Admin can delete any event (except defaults) regardless of bookings"""
        try:
            # Protect default events (IDs 1-3)
            if event_id <= 3:
                return False
            
            events = self.get_events()
            bookings = self.get_bookings()
            
            # Admin has full access - delete event regardless of bookings
            # Remove event completely
            events = [e for e in events if not (e and e.get('EventID') == event_id)]
            
            # Remove ALL bookings for this event (confirmed, cancelled, etc.)
            bookings = [b for b in bookings if not (b and b.get('EventID') == event_id)]
            
            # Save changes
            success1 = self.write_excel(events, self.events_file, ['EventID', 'Title', 'Category', 'Date', 'Venue', 'OrganizerID', 'TotalSeats', 'AvailableSeats', 'Price'])
            success2 = self.write_excel(bookings, self.bookings_file, ['BookingID', 'UserID', 'EventID', 'SeatsBooked', 'BookingDate', 'Status', 'Amount', 'CardLast4'])
            
            return success1 and success2
        except Exception as e:
            print(f"Error in admin_delete_event: {str(e)}")
            return False